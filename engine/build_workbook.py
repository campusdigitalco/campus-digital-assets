"""Build College Money Toolkit v2 (blank sale version + demo version for screenshots)."""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule, FormulaRule
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

NAVY, GREEN, INK, MUTED = "1F4E79", "2E8B57", "1A1A1A", "5F6B7A"
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor="EEF3F8")
KPI_FILL = PatternFill("solid", fgColor="F4F8FB")
F = "Arial"
thin = Side(style="thin", color="D5DCE4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '$#,##0.00;[Red]-$#,##0.00;"-"'
MONEY0 = '$#,##0;[Red]-$#,##0;"-"'
PCT = '0%;-0%;"-"'
DATE = "mmm d, yyyy"

CATS = ["Housing", "Groceries", "Dining out", "Transportation", "Utilities", "Subscriptions",
        "Shopping", "Entertainment", "School", "Health", "Other"]
EXP_ROWS = 500   # rows 5..504
EXP_FIRST, EXP_LAST = 5, 4 + EXP_ROWS


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, size=18, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w


def style(c, fmt=None, inp=False, bold=False, color=INK, align=None):
    c.font = Font(name=F, size=10, bold=bold, color=color)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = INPUT_FILL
    c.alignment = Alignment(horizontal=align, vertical="center")


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def build(demo: bool, path: str):
    wb = Workbook()
    month_start = dt.date(2026, 10, 1)

    # ---------------- START HERE ----------------
    s = wb.active
    s.title = "START HERE"
    title(s, "College Money Toolkit", "Your money, organized. Fill in the yellow cells — everything else calculates for you.")
    header(s, 4, ["Step", "Sheet", "What to do"], [8, 22, 90])
    steps = [
        ("1", "Monthly Budget", "Pick the month you're tracking, enter your income, and set a target for each category."),
        ("2", "Expense Tracker", "Log every purchase: date, what it was, category (dropdown), amount. Budget & Dashboard update automatically."),
        ("3", "Bills", "List recurring bills with the day they're due. Mark Paid / Unpaid each month."),
        ("4", "Savings Goals", "Add goals (spring break, emergency fund, car). See % done and how much to save per month."),
        ("5", "Student Loan Planner", "Enter each loan's balance, rate, and payment to see your payoff date, total interest, and what extra payments save."),
        ("6", "Semester Planner", "List tuition, books, fees, and big costs before they hit. See what's left and how much to set aside monthly."),
        ("7", "Dashboard", "Your month at a glance: income, spending, savings rate, and budget-vs-actual charts."),
    ]
    for i, (a, b, c) in enumerate(steps, 5):
        for j, v in enumerate((a, b, c), 1):
            cell = s.cell(row=i, column=j, value=v)
            style(cell, align="center" if j == 1 else None, bold=(j == 2))
            if i % 2 == 0:
                cell.fill = BAND_FILL
        s.row_dimensions[i].height = 22
    s["A13"] = "Legend"
    s["A13"].font = Font(name=F, bold=True, color=NAVY)
    s["A14"].fill = INPUT_FILL
    s["A14"].border = BOX
    s["B14"] = "Yellow cell = type here"
    s["B15"] = "White cell = calculated automatically (don't type over these)"
    s["B16"] = "Rows marked EXAMPLE show the format — overwrite or delete them."
    s["A18"] = "Works in Microsoft Excel and Google Sheets (File → Import). Educational tool only — not financial advice."
    s["A18"].font = Font(name=F, size=9, italic=True, color=MUTED)
    for r in (14, 15, 16):
        s[f"B{r}"].font = Font(name=F, size=10)
    page(s)

    # ---------------- MONTHLY BUDGET ----------------
    b = wb.create_sheet("Monthly Budget")
    title(b, "Monthly Budget", "Set targets, then watch Actual fill in from the Expense Tracker for the month you pick.")
    b["A3"] = "Month tracking:"
    b["A3"].font = Font(name=F, bold=True)
    b["B3"] = month_start if demo else "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    style(b["B3"], "mmmm yyyy", inp=True, bold=True)
    b["B3"].comment = Comment("Type the 1st of any month (e.g. 9/1/2026) to view that month.", "Toolkit")
    header(b, 5, ["Category", "Monthly target", "Actual spent", "Remaining", "% of target used"], [22, 17, 17, 17, 18])
    targets = dict(zip(CATS, [650, 220, 120, 40, 60, 35, 80, 70, 40, 25, 30]))
    for i, cat in enumerate(CATS, 6):
        b.cell(row=i, column=1, value=cat)
        b.cell(row=i, column=2, value=targets[cat] if demo else (200 if cat == "Groceries" else 0))
        b.cell(row=i, column=3, value=(
            f"=SUMIFS('Expense Tracker'!$D${EXP_FIRST}:$D${EXP_LAST},'Expense Tracker'!$C${EXP_FIRST}:$C${EXP_LAST},A{i},"
            f"'Expense Tracker'!$A${EXP_FIRST}:$A${EXP_LAST},\">=\"&$B$3,'Expense Tracker'!$A${EXP_FIRST}:$A${EXP_LAST},\"<\"&EDATE($B$3,1))"))
        b.cell(row=i, column=4, value=f"=B{i}-C{i}")
        b.cell(row=i, column=5, value=f'=IF(B{i}>0,C{i}/B{i},"")')
        style(b.cell(row=i, column=1), bold=True)
        style(b.cell(row=i, column=2), MONEY, inp=True)
        style(b.cell(row=i, column=3), MONEY)
        style(b.cell(row=i, column=4), MONEY)
        style(b.cell(row=i, column=5), PCT, align="center")
    tr = 6 + len(CATS)  # 17
    b.cell(row=tr, column=1, value="TOTAL")
    for col, L in ((2, "B"), (3, "C"), (4, "D")):
        b.cell(row=tr, column=col, value=f"=SUM({L}6:{L}{tr-1})")
        style(b.cell(row=tr, column=col), MONEY, bold=True)
    b.cell(row=tr, column=5, value=f'=IF(B{tr}>0,C{tr}/B{tr},"")')
    style(b.cell(row=tr, column=5), PCT, bold=True, align="center")
    style(b.cell(row=tr, column=1), bold=True)
    for col in range(1, 6):
        b.cell(row=tr, column=col).fill = BAND_FILL
    # over-budget highlighting
    b.conditional_formatting.add(f"D6:D{tr}", CellIsRule(operator="lessThan", formula=["0"],
                                 font=Font(color="C0392B", bold=True), fill=PatternFill("solid", fgColor="FDECEA")))
    b.conditional_formatting.add(f"E6:E{tr-1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                 end_value=1, color="2E8B57"))

    ir = tr + 2  # 19
    b.cell(row=ir, column=1, value="INCOME").font = Font(name=F, bold=True, color=NAVY, size=12)
    header(b, ir + 1, ["Source", "Monthly amount"])
    incomes = [("Job / paycheck", 900), ("Family support", 400), ("Financial aid / refund", 0), ("Side hustle", 150), ("Other", 0)]
    for k, (src, amt) in enumerate(incomes):
        r = ir + 2 + k
        b.cell(row=r, column=1, value=src)
        b.cell(row=r, column=2, value=amt if demo else (1200 if k == 0 else 0))
        style(b.cell(row=r, column=1), inp=True)
        style(b.cell(row=r, column=2), MONEY, inp=True)
    inc_total = ir + 2 + len(incomes)  # 26
    b.cell(row=inc_total, column=1, value="Total income")
    b.cell(row=inc_total, column=2, value=f"=SUM(B{ir+2}:B{inc_total-1})")
    style(b.cell(row=inc_total, column=1), bold=True)
    style(b.cell(row=inc_total, column=2), MONEY, bold=True)

    sr = inc_total + 2  # 28
    rows = [("Monthly savings goal", 80 if demo else 0, True, MONEY),
            ("Planned left over (income − targets − savings)", f"=B{inc_total}-B{tr}-B{sr}", False, MONEY),
            ("Actual left over so far (income − spent)", f"=B{inc_total}-C{tr}", False, MONEY)]
    for k, (lab, val, inp, fmt) in enumerate(rows):
        r = sr + k
        b.cell(row=r, column=1, value=lab)
        b.cell(row=r, column=2, value=val)
        style(b.cell(row=r, column=1), bold=True)
        style(b.cell(row=r, column=2), fmt, inp=inp, bold=not inp)
    b.column_dimensions["A"].width = 44
    b.freeze_panes = "A6"
    page(b, landscape=False)
    INC = f"'Monthly Budget'!$B${inc_total}"
    SAVE_GOAL = f"'Monthly Budget'!$B${sr}"
    MONTH = "'Monthly Budget'!$B$3"

    # ---------------- EXPENSE TRACKER ----------------
    e = wb.create_sheet("Expense Tracker")
    title(e, "Expense Tracker", "Log every purchase. Category drives your budget — pick from the dropdown.")
    header(e, 4, ["Date", "Description", "Category", "Amount", "Payment method", "Notes"], [14, 32, 18, 14, 18, 30])
    dv = DataValidation(type="list", formula1=f"='Monthly Budget'!$A$6:$A${5+len(CATS)}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Pick a category", error="Choose a category from the list.")
    e.add_data_validation(dv)
    dv.add(f"C{EXP_FIRST}:C{EXP_LAST}")
    pm = DataValidation(type="list", formula1='"Debit card,Credit card,Cash,Venmo,Apple Pay,Meal plan,Other"', allow_blank=True)
    e.add_data_validation(pm)
    pm.add(f"E{EXP_FIRST}:E{EXP_LAST}")
    d = lambda day: dt.date(2026, 10, day)
    sample = [
        (d(1), "October rent", "Housing", 650, "Debit card"), (d(1), "Spotify Student", "Subscriptions", 5.99, "Credit card"),
        (d(2), "Kroger", "Groceries", 58.42, "Debit card"), (d(2), "Chipotle", "Dining out", 12.85, "Apple Pay"),
        (d(3), "Textbook access code", "School", 39.99, "Credit card"), (d(4), "DoorDash — late night", "Dining out", 24.10, "Credit card"),
        (d(4), "Bus pass top-up", "Transportation", 20, "Debit card"), (d(5), "Netflix (split)", "Subscriptions", 7.75, "Venmo"),
        (d(6), "Football game tailgate", "Entertainment", 32, "Venmo"), (d(7), "Electric bill share", "Utilities", 38.5, "Venmo"),
        (d(8), "Starbucks", "Dining out", 6.45, "Apple Pay"), (d(9), "Aldi", "Groceries", 41.17, "Debit card"),
        (d(10), "DoorDash", "Dining out", 19.80, "Credit card"), (d(10), "Target — dorm stuff", "Shopping", 34.60, "Debit card"),
        (d(11), "Movie night", "Entertainment", 14, "Apple Pay"), (d(12), "Internet share", "Utilities", 20, "Venmo"),
        (d(12), "CVS", "Health", 11.29, "Debit card"), (d(13), "Sweetgreen", "Dining out", 14.25, "Apple Pay"),
        (d(14), "Kroger", "Groceries", 47.88, "Debit card"), (d(15), "Uber home from Kilroy's", "Transportation", 13.40, "Credit card"),
        (d(15), "DoorDash", "Dining out", 22.60, "Credit card"), (d(16), "Amazon — phone charger", "Shopping", 15.99, "Credit card"),
        (d(17), "Concert tickets", "Entertainment", 45, "Credit card"), (d(17), "iCloud", "Subscriptions", 2.99, "Credit card"),
        (d(18), "Insomnia Cookies", "Dining out", 9.50, "Apple Pay"),
    ]
    rows_to_write = sample if demo else [("=TODAY()", "EXAMPLE — Kroger groceries (overwrite me)", "Groceries", 45.00, "Debit card")]
    for k, row in enumerate(rows_to_write):
        r = EXP_FIRST + k
        for j, v in enumerate(row, 1):
            e.cell(row=r, column=j, value=v)
    for r in range(EXP_FIRST, EXP_LAST + 1):
        style(e.cell(row=r, column=1), DATE, inp=True)
        for j in (2, 3, 5, 6):
            style(e.cell(row=r, column=j), inp=True)
        style(e.cell(row=r, column=4), MONEY, inp=True)
    e["H4"] = "This month"
    e["H4"].font = Font(name=F, bold=True, color="FFFFFF")
    e["H4"].fill = HEAD_FILL
    e["H5"] = "Transactions"
    e["I5"] = (f"=COUNTIFS(A{EXP_FIRST}:A{EXP_LAST},\">=\"&{MONTH},A{EXP_FIRST}:A{EXP_LAST},\"<\"&EDATE({MONTH},1))")
    e["H6"] = "Total spent"
    e["I6"] = (f"=SUMIFS(D{EXP_FIRST}:D{EXP_LAST},A{EXP_FIRST}:A{EXP_LAST},\">=\"&{MONTH},A{EXP_FIRST}:A{EXP_LAST},\"<\"&EDATE({MONTH},1))")
    e["H7"] = "Biggest purchase"
    e["I7"] = (f"=_xlfn.MAXIFS(D{EXP_FIRST}:D{EXP_LAST},A{EXP_FIRST}:A{EXP_LAST},\">=\"&{MONTH},A{EXP_FIRST}:A{EXP_LAST},\"<\"&EDATE({MONTH},1))")
    for r in (5, 6, 7):
        style(e[f"H{r}"], bold=True)
        style(e[f"I{r}"], MONEY if r > 5 else "0")
    e.column_dimensions["H"].width = 18
    e.column_dimensions["I"].width = 14
    e.freeze_panes = "A5"
    e.print_area = f"A1:I{EXP_FIRST + max(len(rows_to_write), 10) + 2}"
    page(e)

    # ---------------- BILLS ----------------
    bl = wb.create_sheet("Bills")
    title(bl, "Bills", "Recurring bills by due day. Mark Paid / Unpaid each month — unpaid totals show on the Dashboard.")
    header(bl, 4, ["Due day", "Bill", "Amount", "Autopay?", "Status", "Notes"], [10, 28, 14, 12, 12, 30])
    sv = DataValidation(type="list", formula1='"Paid,Unpaid"', allow_blank=True)
    yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    bl.add_data_validation(sv)
    bl.add_data_validation(yn)
    sv.add("E5:E54")
    yn.add("D5:D54")
    bills = [(1, "Rent", 650, "No", "Paid"), (5, "Phone (family plan share)", 30, "Yes", "Paid"),
             (12, "Internet share", 20, "No", "Paid"), (15, "Car insurance", 68, "Yes", "Unpaid"),
             (20, "Electric share", 38.5, "No", "Paid"), (25, "Gym membership", 15, "Yes", "Unpaid")]
    btw = bills if demo else [(1, "EXAMPLE — Rent", 650, "No", "Unpaid")]
    for k, row in enumerate(btw):
        for j, v in enumerate(row, 1):
            bl.cell(row=5 + k, column=j, value=v)
    for r in range(5, 55):
        style(bl.cell(row=r, column=1), "0", inp=True, align="center")
        style(bl.cell(row=r, column=2), inp=True)
        style(bl.cell(row=r, column=3), MONEY, inp=True)
        for j in (4, 5, 6):
            style(bl.cell(row=r, column=j), inp=True, align="center" if j < 6 else None)
    bl.conditional_formatting.add("E5:E54", CellIsRule(operator="equal", formula=['"Unpaid"'],
                                  font=Font(color="C0392B", bold=True), fill=PatternFill("solid", fgColor="FDECEA")))
    bl.conditional_formatting.add("E5:E54", CellIsRule(operator="equal", formula=['"Paid"'],
                                  font=Font(color="1E7B47", bold=True), fill=PatternFill("solid", fgColor="E6F4EC")))
    bl["H4"], bl["H5"], bl["H6"] = "Summary", "Total monthly bills", "Still unpaid"
    bl["I5"], bl["I6"] = "=SUM(C5:C54)", '=SUMIF(E5:E54,"Unpaid",C5:C54)'
    bl["H4"].font = Font(name=F, bold=True, color="FFFFFF")
    bl["H4"].fill = HEAD_FILL
    for r in (5, 6):
        style(bl[f"H{r}"], bold=True)
        style(bl[f"I{r}"], MONEY, bold=True)
    bl.column_dimensions["H"].width = 20
    bl.column_dimensions["I"].width = 14
    bl.freeze_panes = "A5"
    bl.print_area = "A1:I16"
    page(bl)

    # ---------------- SAVINGS GOALS ----------------
    sg = wb.create_sheet("Savings Goals")
    title(sg, "Savings Goals", "Name it, set a target and date — see your progress and exactly what to save each month.")
    header(sg, 4, ["Goal", "Target", "Saved so far", "Target date", "% complete", "Still need", "Months left", "Save per month"],
           [26, 13, 14, 15, 16, 13, 12, 16])
    goals = [("Emergency fund", 1000, 420, dt.date(2027, 5, 1)), ("Spring break", 600, 260, dt.date(2027, 3, 1)),
             ("New laptop", 1200, 150, dt.date(2027, 8, 1)), ("Summer apartment deposit", 500, 0, dt.date(2027, 4, 15))]
    gtw = goals if demo else [("EXAMPLE — Spring break", 600, 100, dt.date(2027, 3, 1))]
    for k, row in enumerate(gtw):
        for j, v in enumerate(row, 1):
            sg.cell(row=5 + k, column=j, value=v)
    for r in range(5, 25):
        style(sg.cell(row=r, column=1), inp=True, bold=True)
        style(sg.cell(row=r, column=2), MONEY0, inp=True)
        style(sg.cell(row=r, column=3), MONEY0, inp=True)
        style(sg.cell(row=r, column=4), DATE, inp=True)
        sg.cell(row=r, column=5, value=f'=IF(B{r}>0,MIN(C{r}/B{r},1),"")')
        sg.cell(row=r, column=6, value=f'=IF(B{r}>0,MAX(B{r}-C{r},0),"")')
        sg.cell(row=r, column=7, value=f'=IF(AND(B{r}>0,ISNUMBER(D{r})),IF(D{r}>TODAY(),MAX((YEAR(D{r})-YEAR(TODAY()))*12+MONTH(D{r})-MONTH(TODAY()),1),0),"")')
        sg.cell(row=r, column=8, value=f'=IF(G{r}="","",IF(F{r}=0,0,IF(G{r}=0,"Date passed",F{r}/G{r})))')
        style(sg.cell(row=r, column=5), PCT, align="center")
        style(sg.cell(row=r, column=6), MONEY0)
        style(sg.cell(row=r, column=7), "0", align="center")
        style(sg.cell(row=r, column=8), MONEY, bold=True)
    sg.conditional_formatting.add("E5:E24", DataBarRule(start_type="num", start_value=0, end_type="num",
                                  end_value=1, color="2E8B57"))
    sg["A26"], sg["B26"], sg["C26"], sg["F26"], sg["H26"] = "TOTAL", "=SUM(B5:B24)", "=SUM(C5:C24)", "=SUM(F5:F24)", "=SUM(H5:H24)"
    sg["E26"] = '=IF(B26>0,C26/B26,"")'
    for col, fmt in (("A", None), ("B", MONEY0), ("C", MONEY0), ("E", PCT), ("F", MONEY0), ("H", MONEY)):
        style(sg[f"{col}26"], fmt, bold=True)
        sg[f"{col}26"].fill = BAND_FILL
    sg["A28"] = "Months left counts from today. Save per month = still need ÷ months left."
    sg["A28"].font = Font(name=F, size=9, italic=True, color=MUTED)
    sg.freeze_panes = "A5"
    sg.print_area = "A1:H28"
    page(sg)

    # ---------------- STUDENT LOAN PLANNER ----------------
    ln = wb.create_sheet("Student Loan Planner")
    title(ln, "Student Loan Planner", "See your payoff date, total interest — and how much an extra payment saves you.")
    header(ln, 4, ["Loan", "Balance", "Interest rate", "Monthly payment", "Extra / month", "Months to pay off",
                   "Payoff date", "Total interest", "Months w/ extra", "Interest saved w/ extra"],
           [24, 13, 12, 14, 13, 13, 14, 14, 13, 16])
    loans = [("Federal Direct Sub", 5500, 0.0639, 60, 25), ("Federal Direct Unsub", 7000, 0.0639, 80, 25),
             ("Private loan", 4000, 0.089, 75, 0)]
    ltw = loans if demo else [("EXAMPLE — Direct Unsubsidized", 5500, 0.0639, 65, 25)]
    for k, row in enumerate(ltw):
        for j, v in enumerate(row, 1):
            ln.cell(row=5 + k, column=j, value=v)
    for r in range(5, 25):
        style(ln.cell(row=r, column=1), inp=True, bold=True)
        style(ln.cell(row=r, column=2), MONEY0, inp=True)
        style(ln.cell(row=r, column=3), '0.00%;-0.00%;"-"', inp=True, align="center")
        style(ln.cell(row=r, column=4), MONEY0, inp=True)
        style(ln.cell(row=r, column=5), MONEY0, inp=True)
        ok = f"AND(B{r}>0,D{r}>0)"
        low = f"D{r}<=B{r}*C{r}/12"
        months = f"IF(C{r}=0,ROUNDUP(B{r}/D{r},0),ROUNDUP(NPER(C{r}/12,-D{r},B{r}),0))"
        ln.cell(row=r, column=6, value=f'=IF(NOT({ok}),"",IF({low},"Payment too low",{months}))')
        ln.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),EDATE(TODAY(),F{r}),"")')
        tot_int = f"IF(C{r}=0,0,NPER(C{r}/12,-D{r},B{r})*D{r}-B{r})"
        ln.cell(row=r, column=8, value=f'=IF(ISNUMBER(F{r}),{tot_int},"")')
        pe = f"(D{r}+E{r})"
        months_x = f"IF(C{r}=0,ROUNDUP(B{r}/{pe},0),ROUNDUP(NPER(C{r}/12,-{pe},B{r}),0))"
        ln.cell(row=r, column=9, value=f'=IF(AND(ISNUMBER(F{r}),E{r}>0),{months_x},"")')
        int_x = f"IF(C{r}=0,0,NPER(C{r}/12,-{pe},B{r})*{pe}-B{r})"
        ln.cell(row=r, column=10, value=f'=IF(ISNUMBER(I{r}),H{r}-({int_x}),"")')
        style(ln.cell(row=r, column=6), "0", align="center")
        style(ln.cell(row=r, column=7), "mmm yyyy", align="center")
        style(ln.cell(row=r, column=8), MONEY0)
        style(ln.cell(row=r, column=9), "0", align="center")
        style(ln.cell(row=r, column=10), MONEY0, bold=True, color="1E7B47")
    ln["A26"] = "TOTAL"
    for col, fmt in (("B", MONEY0), ("D", MONEY0), ("E", MONEY0), ("H", MONEY0), ("J", MONEY0)):
        ln[f"{col}26"] = f"=SUM({col}5:{col}24)"
        style(ln[f"{col}26"], fmt, bold=True)
        ln[f"{col}26"].fill = BAND_FILL
    style(ln["A26"], bold=True)
    ln["A26"].fill = BAND_FILL
    ln["A28"] = ("Enter the rate as a percent (6.39%). Estimates assume a fixed rate and payments starting now; "
                 "check exact figures with your loan servicer (studentaid.gov for federal loans).")
    ln["A28"].font = Font(name=F, size=9, italic=True, color=MUTED)
    ln.freeze_panes = "A5"
    ln.print_area = "A1:J28"
    page(ln)

    # ---------------- SEMESTER PLANNER ----------------
    sp = wb.create_sheet("Semester Planner")
    title(sp, "Semester Planner", "Map out the big costs before they hit, so nothing sneaks up on you.")
    sp["A3"], sp["B3"] = "Semester:", ("Spring 2027" if demo else "Spring 2027")
    sp["D3"], sp["E3"] = "Months left:", (5 if demo else 4)
    sp["A3"].font = sp["D3"].font = Font(name=F, bold=True)
    style(sp["B3"], inp=True, bold=True)
    style(sp["E3"], "0", inp=True, bold=True, align="center")
    header(sp, 5, ["Item", "Category", "Estimated cost", "Due date", "Amount paid", "Still owe", "Status"],
           [30, 16, 15, 15, 14, 14, 14])
    cat_dv = DataValidation(type="list", formula1='"Tuition,Housing,Books,Fees,Supplies,Travel,Meal plan,Other"', allow_blank=True)
    sp.add_data_validation(cat_dv)
    cat_dv.add("B6:B35")
    items = [("Tuition balance after aid", "Tuition", 1850, dt.date(2027, 1, 10), 0),
             ("Textbooks & access codes", "Books", 320, dt.date(2027, 1, 15), 0),
             ("Parking permit", "Fees", 185, dt.date(2027, 1, 12), 185),
             ("Lab / course fees", "Fees", 120, dt.date(2027, 1, 20), 0),
             ("Spring break trip", "Travel", 600, dt.date(2027, 3, 13), 260),
             ("Flights home (Thanksgiving)", "Travel", 240, dt.date(2026, 11, 20), 240),
             ("Laptop repair / supplies", "Supplies", 90, dt.date(2027, 2, 1), 0),
             ("Formal / event dues", "Other", 150, dt.date(2027, 2, 15), 50)]
    itw = items if demo else [("EXAMPLE — Textbooks", "Books", 300, dt.date(2027, 1, 15), 0)]
    for k, row in enumerate(itw):
        for j, v in enumerate(row, 1):
            sp.cell(row=6 + k, column=j, value=v)
    for r in range(6, 36):
        style(sp.cell(row=r, column=1), inp=True, bold=True)
        style(sp.cell(row=r, column=2), inp=True, align="center")
        style(sp.cell(row=r, column=3), MONEY0, inp=True)
        style(sp.cell(row=r, column=4), DATE, inp=True)
        style(sp.cell(row=r, column=5), MONEY0, inp=True)
        sp.cell(row=r, column=6, value=f'=IF(C{r}="","",MAX(C{r}-E{r},0))')
        sp.cell(row=r, column=7, value=f'=IF(C{r}="","",IF(F{r}=0,"Paid",IF(AND(ISNUMBER(D{r}),D{r}<TODAY()),"Overdue",IF(E{r}>0,"Partial","Upcoming"))))')
        style(sp.cell(row=r, column=6), MONEY0)
        style(sp.cell(row=r, column=7), align="center", bold=True)
    sp.conditional_formatting.add("G6:G35", CellIsRule(operator="equal", formula=['"Paid"'], font=Font(color="1E7B47", bold=True),
                                  fill=PatternFill("solid", fgColor="E6F4EC")))
    sp.conditional_formatting.add("G6:G35", CellIsRule(operator="equal", formula=['"Overdue"'], font=Font(color="C0392B", bold=True),
                                  fill=PatternFill("solid", fgColor="FDECEA")))
    sp.conditional_formatting.add("G6:G35", CellIsRule(operator="equal", formula=['"Partial"'], font=Font(color="9A6B00", bold=True),
                                  fill=PatternFill("solid", fgColor="FFF6D5")))
    sp["A37"], sp["C37"], sp["E37"], sp["F37"] = "TOTAL", "=SUM(C6:C35)", "=SUM(E6:E35)", "=SUM(F6:F35)"
    for col, fmt in (("A", None), ("C", MONEY0), ("E", MONEY0), ("F", MONEY0)):
        style(sp[f"{col}37"], fmt, bold=True)
        sp[f"{col}37"].fill = BAND_FILL
    sp["A39"] = "Set aside per month to cover it all:"
    sp["A39"].font = Font(name=F, bold=True, size=12, color=NAVY)
    sp["C39"] = '=IF(E3>0,F37/E3,F37)'
    style(sp["C39"], MONEY, bold=True, color="1E7B47")
    sp["C39"].font = Font(name=F, bold=True, size=14, color="1E7B47")
    sp.freeze_panes = "A6"
    sp.print_area = "A1:G39"
    page(sp)

    # ---------------- DASHBOARD ----------------
    db = wb.create_sheet("Dashboard", 1)
    title(db, "Your Money — At a Glance", "Updates automatically from every other sheet.")
    db["A3"] = "Month:"
    db["B3"] = f"={MONTH}"
    db["A3"].font = Font(name=F, bold=True)
    db["B3"].number_format = "mmmm yyyy"
    db["B3"].font = Font(name=F, bold=True, color=NAVY, size=12)
    header(db, 5, ["Metric", "Amount", "What it means"], [30, 16, 46])
    kpis = [
        ("Income this month", f"={INC}", MONEY, "Total from the Income section of Monthly Budget"),
        ("Spent this month", f"='Monthly Budget'!C{tr}", MONEY, "Everything logged in Expense Tracker this month"),
        ("Left to spend (budget)", f"='Monthly Budget'!D{tr}", MONEY, "Budget targets minus actual spending"),
        ("Cash left after spending", "=B6-B7", MONEY, "Income minus what you've spent"),
        ("Savings rate", '=IF(B6>0,MAX(B9,0)/B6,"")', PCT, "Share of income you're keeping this month"),
        ("Monthly savings goal", f"={SAVE_GOAL}", MONEY, "Your target from Monthly Budget"),
        ("On track for savings goal?", '=IF(B11=0,"Set a goal",IF(B9>=B11,"Yes ✓","Not yet"))', None, "Cash left vs. savings goal"),
        ("Bills still unpaid", "=Bills!I6", MONEY, "Bills marked Unpaid"),
        ("Savings goals progress", "='Savings Goals'!E26", PCT, "Saved so far ÷ total of all goals"),
        ("Total student loan balance", "='Student Loan Planner'!B26", MONEY0, "Sum of all loan balances"),
        ("Semester costs still owed", "='Semester Planner'!F37", MONEY0, "From Semester Planner"),
    ]
    for k, (lab, f, fmt, note) in enumerate(kpis):
        r = 6 + k
        db.cell(row=r, column=1, value=lab)
        db.cell(row=r, column=2, value=f)
        db.cell(row=r, column=3, value=note)
        style(db.cell(row=r, column=1), bold=True)
        style(db.cell(row=r, column=2), fmt, bold=True, color=NAVY, align="right")
        style(db.cell(row=r, column=3), color=MUTED)
        db.cell(row=r, column=2).font = Font(name=F, size=12, bold=True, color=NAVY)
        if k % 2 == 0:
            for c in range(1, 4):
                db.cell(row=r, column=c).fill = KPI_FILL
        db.row_dimensions[r].height = 22
    db.conditional_formatting.add("B9", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C0392B", bold=True)))
    db.conditional_formatting.add("B8", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C0392B", bold=True)))

    ch = BarChart()
    ch.type = "bar"
    ch.title = "Budget vs. actual by category"
    ch.style = 10
    data = Reference(b, min_col=2, max_col=3, min_row=5, max_row=5 + len(CATS))
    cats_ref = Reference(b, min_col=1, min_row=6, max_row=5 + len(CATS))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.series[0].graphicalProperties.solidFill = "B9C9DA"
    ch.series[0].graphicalProperties.line.solidFill = "B9C9DA"
    ch.series[1].graphicalProperties.solidFill = NAVY
    ch.series[1].graphicalProperties.line.solidFill = NAVY
    ch.y_axis.numFmt = "$#,##0"
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.height, ch.width = 10, 17
    ch.legend.position = "b"
    db.add_chart(ch, "E5")

    dn = DoughnutChart()
    dn.title = "Where your money went"
    dn.add_data(Reference(b, min_col=3, min_row=5, max_row=5 + len(CATS)), titles_from_data=True)
    dn.set_categories(cats_ref)
    dn.holeSize = 55
    dn.height, dn.width = 10, 17
    dn.legend.position = "r"
    from openpyxl.chart.series import DataPoint
    pal = ["1F4E79", "2E8B57", "E07A5F", "F2CC8F", "81B29A", "3D5A80", "98C1D9", "C0392B", "8D6A9F", "6C757D", "B9C9DA"]
    for i, col in enumerate(pal):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = col
        pt.graphicalProperties.line.solidFill = "FFFFFF"
        dn.series[0].dPt.append(pt)
    dn.height, dn.width = 9.5, 16.5
    db.add_chart(dn, "A19")
    for col in "EFGHIJKL":
        db.column_dimensions[col].width = 11
    db.print_area = "A1:M38"
    page(db)

    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    build(demo=False, path=f"{out}/College_Money_Toolkit_v2.xlsx")
    build(demo=True, path=f"{out}/College_Money_Toolkit_DEMO.xlsx")
    print("built")
